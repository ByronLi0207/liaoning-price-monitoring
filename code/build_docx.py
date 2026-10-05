#!/usr/bin/env python3
"""Native Word from frozen Markdown; editable OMML, tables, and real plot files."""
from pathlib import Path
import re,json,hashlib,copy,argparse,zipfile
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START,WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
M='http://schemas.openxmlformats.org/officeDocument/2006/math'
def clean_metadata(dest):
 from lxml import etree
 with zipfile.ZipFile(dest) as z:parts={n:z.read(n) for n in z.namelist()}
 for name in ['docProps/core.xml','docProps/app.xml']:
  tree=etree.fromstring(parts[name])
  for e in list(tree):
   if etree.QName(e).localname in ['created','modified','lastPrinted','revision','Pages','Words','Characters','CharactersWithSpaces','TotalTime','Application','AppVersion','Company','Manager','Template']:
    tree.remove(e)
  parts[name]=etree.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
 tmp=dest.with_suffix('.tmp')
 with zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as z:
  for name,data in parts.items():z.writestr(name,data)
 tmp.replace(dest)
def el(tag,*children,**attrs):
 n=OxmlElement('m:'+tag)
 for k,v in attrs.items():n.set(qn('m:'+k),str(v))
 for c in children:n.append(c)
 return n
def text(s):
 r=el('r');t=el('t');t.text=s;t.set(qn('xml:space'),'preserve');r.append(t);return r
def wrap(tag,items):return el(tag,*items)
class MathParser:
 symbols={'times':'×','cdot':'·','le':'≤','leq':'≤','ge':'≥','geq':'≥','in':'∈','sum':'∑','tau':'τ','Delta':'Δ','bar':'¯','varnothing':'∅','ldots':'…','ldotp':'.','mid':'|','to':'→','pm':'±'}
 def __init__(self,s):self.s=s;self.i=0
 def atom(self):
  while self.i<len(self.s) and self.s[self.i].isspace():self.i+=1
  if self.i>=len(self.s):return text('')
  c=self.s[self.i];self.i+=1
  if c=='{':return self.seq('}')
  if c=='\\':
   m=re.match(r'[A-Za-z]+',self.s[self.i:])
   if m:cmd=m.group();self.i+=len(cmd)
   else:cmd=self.s[self.i:self.i+1];self.i+=1
   if cmd in ['left','right']:return self.atom()
   if cmd=='frac':
    a,b=self.items(self.atom()),self.items(self.atom());return el('f',wrap('num',a),wrap('den',b))
   if cmd=='sqrt':
    pr=el('radPr',el('degHide',val='1'));return el('rad',pr,el('deg'),wrap('e',self.items(self.atom())))
   if cmd in ['overline','bar']:return el('bar',el('barPr',el('pos',val='top')),wrap('e',self.items(self.atom())))
   if cmd in ['widehat','hat']:return el('acc',el('accPr',el('chr',val='̂')),wrap('e',self.items(self.atom())))
   if cmd in ['mathrm','text','operatorname','mathbf']:return self.atom()
   if cmd in ['log','exp','min','max']:return text(cmd)
   if cmd=='ell':return text('ℓ')
   if cmd in [',',';','quad','qquad',' ']:return text(' ')
   if cmd=='!':return text('')
   if cmd in ['{','}','_','%']:return text(cmd)
   if cmd in self.symbols:return text(self.symbols[cmd])
   raise ValueError('Unsupported TeX command '+cmd+' in '+self.s)
  return text('−' if c=='-' else c)
 def items(self,a):return a if isinstance(a,list) else [a]
 def seq(self,end=None):
  out=[]
  while self.i<len(self.s):
   if end and self.s[self.i]==end:self.i+=1;break
   if self.s[self.i]=='&':self.i+=1;continue
   a=self.atom();sub=sup=None
   while self.i<len(self.s) and self.s[self.i] in '_^':
    op=self.s[self.i];self.i+=1;b=self.items(self.atom())
    if op=='_':sub=b
    else:sup=b
   if sub is not None and sup is not None:a=el('sSubSup',wrap('e',self.items(a)),wrap('sub',sub),wrap('sup',sup))
   elif sub is not None:a=el('sSub',wrap('e',self.items(a)),wrap('sub',sub))
   elif sup is not None:a=el('sSup',wrap('e',self.items(a)),wrap('sup',sup))
   out.extend(self.items(a))
  return out
def omml(tex,display=False):
 tex=re.sub(r'\\tag\{[^{}]*\}','',tex).strip()
 tex=tex.replace('\\begin{aligned}','').replace('\\end{aligned}','')
 if '\\\\' in tex:
  rows=tex.split('\\\\');node=el('oMath',el('eqArr',*[wrap('e',MathParser(r).seq()) for r in rows]))
 else:node=el('oMath',*MathParser(tex).seq())
 return el('oMathPara',node) if display else node
def font(run,size=None):
 run.font.name='Liberation Serif';run.font.color.rgb=RGBColor(0,0,0)
 if size:run.font.size=Pt(size)
 rf=run._r.get_or_add_rPr();fonts=rf.find(qn('w:rFonts'))
 if fonts is None:fonts=OxmlElement('w:rFonts');rf.insert(0,fonts)
 for k in ['ascii','hAnsi','eastAsia']:fonts.set(qn('w:'+k),'Liberation Serif')
def add_text(p,s):
 pattern=r'(\\\([\s\S]*?\\\)|\$[^$\n]+\$|\*\*[^*]+\*\*|(?<!\*)\*[^*]+\*(?!\*)|\[[^\]]+\]\([^)]+\))'
 for part in re.split(pattern,s):
  if not part:continue
  if part.startswith('\\('):p._p.append(omml(part[2:-2]));continue
  if part.startswith('$'):p._p.append(omml(part[1:-1]));continue
  m=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',part)
  if m:
   hl=OxmlElement('w:hyperlink');rid=p.part.relate_to(m[2],'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True);hl.set(qn('r:id'),rid)
   rr=OxmlElement('w:r');tt=OxmlElement('w:t');tt.text=m[1];rr.append(tt);hl.append(rr);p._p.append(hl);continue
  bold=part.startswith('**') and part.endswith('**');italic=not bold and part.startswith('*') and part.endswith('*')
  run=p.add_run(part[2:-2] if bold else part[1:-1] if italic else part);font(run);run.bold=bold;run.italic=italic
def section(doc,landscape):
 s=doc.add_section(WD_SECTION_START.NEW_PAGE)
 s.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
 s.page_width=Cm(27.94 if landscape else 21.59);s.page_height=Cm(21.59 if landscape else 27.94)
 s.top_margin=s.bottom_margin=Cm(1.8);s.left_margin=s.right_margin=Cm(1.8);return s
def format_table(t,widths,eq=False):
 t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
 pr=t._tbl.tblPr;b=OxmlElement('w:tblBorders')
 for name in ['top','left','bottom','right','insideH','insideV']:
  q=OxmlElement('w:'+name);q.set(qn('w:val'),'nil' if eq else 'single');q.set(qn('w:sz'),'4');q.set(qn('w:color'),'D9D9D9');b.append(q)
 pr.append(b)
 for i,w in enumerate(widths):t.columns[i].width=Cm(w)
 for ri,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
  if ri==0 and not eq:trpr.append(OxmlElement('w:tblHeader'))
  for i,cell in enumerate(row.cells):
   cell.width=Cm(widths[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcpr=cell._tc.get_or_add_tcPr();pad=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'65');e.set(qn('w:type'),'dxa');pad.append(e)
   tcpr.append(pad)
   if ri==0 and not eq:
    fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E9EEF3');tcpr.append(fill)
   for p in cell.paragraphs:
    p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.05
    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if i==0 else WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:font(run,9.5);run.bold=ri==0 and not eq
def make(md,dest,role):
 doc=Document();s=doc.sections[0];s.page_width=Cm(21.59);s.page_height=Cm(27.94);s.top_margin=s.bottom_margin=Cm(1.8);s.left_margin=s.right_margin=Cm(1.8)
 for name,size in [('Normal',11),('Title',16),('Heading 1',13),('Heading 2',12),('Heading 3',11),('Caption',9.5)]:
  sty=doc.styles[name];sty.font.name='Liberation Serif';sty.font.size=Pt(size);sty.font.color.rgb=RGBColor(0,0,0);sty.paragraph_format.space_after=Pt(6);sty.paragraph_format.line_spacing=1.22
  if name.startswith('Heading'):sty.paragraph_format.keep_with_next=True;sty.paragraph_format.space_before=Pt(10)
 footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),' PAGE ');footer._p.append(fld)
 chunks=re.split(r'\n\s*\n',md.strip());pending_caption=None;return_portrait=False;source_tables=[];eqs=[];figs=[];refs=False
 for c in chunks:
  if not c.strip():continue
  if c.startswith('**Table '):pending_caption=c;continue
  if return_portrait and not c.startswith('*Note:'):
   section(doc,False);return_portrait=False
  if c.startswith('|'):
   lines=[l for l in c.splitlines() if l.startswith('|')];rows=[[v.strip() for v in l.strip().strip('|').split('|')] for l in lines if not re.fullmatch(r'[\s|:\-]+',l)]
   cap=pending_caption or ''
   n=len(rows[0]);assert all(len(r)==n for r in rows)
   wide=n>=6
   if wide:section(doc,True)
   if pending_caption:
    p=doc.add_paragraph();add_text(p,pending_caption);p.paragraph_format.keep_with_next=True;pending_caption=None
   t=doc.add_table(rows=len(rows),cols=n)
   for ri,row in enumerate(rows):
    for ci,v in enumerate(row):add_text(t.cell(ri,ci).paragraphs[0],v)
   widths={3:[3.0,4.1,10.7],5:[2.2,2.0,4.5,4.5,4.5],6:[2.5,3.7,2.9,6.6,2.3,5.6],9:[1.3,1.5,2.1,2.4,2.5,3.4,2.2,4.4,2.6]}.get(n,[24/n]*n)
   if role in ['supplement_methods','supplement_validation']:
    widths={3:[3.4,7.2,7.2],5:[3.4,3.5,4.4,2.5,4.0]}.get(n,widths)
    if role == 'supplement_validation' and n == 4:widths=[4.2,4.2,4.7,4.7]
   if role=='manuscript' and cap.startswith('**Table 5.'):widths=[3.6,3.3,3.4,4.1,3.4]
   if role=='manuscript' and cap.startswith('**Table 6.'):widths=[4.5,2.1,5.6,5.6]
   if role=='manuscript' and cap.startswith('**Table A1.'):widths=[2.8,3.0,5.0,2.5,4.5]
   format_table(t,widths)
   if role == 'supplement_validation':
    for row in t.rows:
     for cell in row.cells:
      for p in cell.paragraphs:
       p.paragraph_format.left_indent=Cm(0);p.paragraph_format.right_indent=Cm(0);p.paragraph_format.first_line_indent=Cm(0)
   source_tables.append(rows);return_portrait=wide;continue
  if pending_caption:
   p=doc.add_paragraph();add_text(p,pending_caption);pending_caption=None
  if c.startswith('$$'):
   tx=c.strip()[2:-2].strip();tag=re.search(r'\\tag\{(\d+)\}',tx);num=tag[1] if tag else ''
   t=doc.add_table(rows=1,cols=2);format_table(t,[16.3,1.4],True)
   p=t.cell(0,0).paragraphs[0];p._p.append(omml(tx,True));p.alignment=WD_ALIGN_PARAGRAPH.CENTER
   t.cell(0,1).paragraphs[0].text='('+num+')';t.cell(0,1).paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.RIGHT;eqs.append({'number':num,'latex':tx});continue
  f=re.fullmatch(r'<!--\s*FIGURE:(S?\d+)\s*-->',c)
  if f:
   fn=ROOT/'figures'/('Figure_'+f[1]+'.png');assert fn.exists()
   p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(fn),width=Cm(17.2));figs.append(str(fn));continue
  hm=re.match(r'^(#{1,4}) (.+)$',c)
  if hm:
   lev=len(hm[1]);title=hm[2]
   if lev==1:p=doc.add_paragraph(title,'Title')
   else:p=doc.add_paragraph(title,'Heading '+str(min(lev-1,3)))
   if title=='References':refs=True
   continue
  if c.startswith('```'):continue
  for line in c.splitlines() if c.startswith('- ') else [c]:
   p=doc.add_paragraph(style='List Bullet' if line.startswith('- ') else None);add_text(p,line[2:] if line.startswith('- ') else line)
   p.paragraph_format.widow_control=True
   if line.startswith('Figure '):
    p.paragraph_format.keep_together=True;p.paragraph_format.line_spacing=1.1
    for run in p.runs:font(run,9.5)
   if refs and not line.startswith('#'):
    p.paragraph_format.left_indent=Cm(.5);p.paragraph_format.first_line_indent=Cm(-.5)
    p.paragraph_format.space_after=Pt(3)
  if return_portrait and c.startswith('*Note:'):section(doc,False);return_portrait=False
 if role == 'project_brief':
  # The supplied OFL font covers the Chinese brief in Word and PDF exports.
  family = 'Agri Serif SC'
  for sty in doc.styles:
   if not hasattr(sty, 'font'):continue
   sty.font.name = family
   rf = sty.element.get_or_add_rPr().find(qn('w:rFonts'))
   if rf is None:
    rf = OxmlElement('w:rFonts');sty.element.get_or_add_rPr().insert(0, rf)
   for key in ['ascii', 'hAnsi', 'eastAsia', 'cs']:rf.set(qn('w:'+key), family)
  for run in doc.element.iter(qn('w:r')):
   pr = run.find(qn('w:rPr'))
   if pr is None:pr = OxmlElement('w:rPr');run.insert(0, pr)
   rf = pr.find(qn('w:rFonts'))
   if rf is None:rf = OxmlElement('w:rFonts');pr.insert(0, rf)
   for key in ['ascii', 'hAnsi', 'eastAsia', 'cs']:rf.set(qn('w:'+key), family)
 doc.core_properties.author='';doc.core_properties.last_modified_by='';doc.core_properties.comments=''
 # Remove the default template's decorative Title borders in native presentation.
 for sty in doc.styles:
  for b in list(sty.element.findall('.//'+qn('w:pBdr'))):b.getparent().remove(b)
 for p in doc.paragraphs:
  for b in list(p._p.findall('.//'+qn('w:pBdr'))):b.getparent().remove(b)
 dest.parent.mkdir(parents=True,exist_ok=True);doc.save(dest);clean_metadata(dest)
 index={'source_md_sha256':hashlib.sha256(md.encode()).hexdigest(),'docx_sha256':hashlib.sha256(dest.read_bytes()).hexdigest(),'source_tables':source_tables,'display_equations':eqs,'figures':figs,'font_requested':'Agri Serif SC' if role == 'project_brief' else 'Liberation Serif','role':role}
 (dest.parent/(dest.stem+'_structure.json')).write_text(json.dumps(index,ensure_ascii=False,indent=2)+'\n', encoding='utf-8')
 print(json.dumps({'file':str(dest),'bytes':dest.stat().st_size,'sha256':index['docx_sha256'],'tables':len(source_tables),'display_equations':len(eqs),'images':len(figs)}))
if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('--role',choices=['manuscript','supplement','supplement_methods','supplement_validation','project_brief'],default='manuscript');a=p.parse_args()
 f=ROOT/'reports'/(a.role+'.md');make(f.read_text(encoding='utf-8'),ROOT/'native'/(a.role+'.docx'),a.role)
