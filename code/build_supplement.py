#!/usr/bin/env python3
"""Reversible English labels and complete keyed column panels; no scientific changes."""
from pathlib import Path
import re,json,hashlib,copy
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from build_docx import add_text,font,format_table,omml,clean_metadata
ROOT=Path(__file__).resolve().parents[1]
sha=lambda b:hashlib.sha256(b).hexdigest()
src=(ROOT/'reports/supplement.md').read_text(encoding='utf-8')
assert sha(src.encode())=='6f16df36f42e45528a34505007d4e27c949e0fba3e1ac66a6e412b2e067d5475'
dp=ROOT/'inputs/presentation_dictionary.json'
d=json.loads(dp.read_text(encoding='utf-8'));mapping={x['zh']:x['en'] for x in d['entries']};mapping.update({'、':'; ','\u3000':' '})
pat=re.compile('|'.join(re.escape(x) for x in sorted(mapping,key=len,reverse=True)))
out=[];log=[];end=0;at=0
for m in pat.finditer(src):
 part=src[end:m.start()];out.append(part);at+=len(part);replacement=mapping[m.group()]
 log.append({'source_offset':m.start(),'presentation_offset':at,'original':m.group(),'replacement':replacement})
 out.append(replacement);at+=len(replacement);end=m.end()
out.append(src[end:]);present=''.join(out)
reverse=present
for r in reversed(log):
 a=r['presentation_offset'];assert reverse[a:a+len(r['replacement'])]==r['replacement'];reverse=reverse[:a]+r['original']+reverse[a+len(r['replacement']):]
assert reverse==src
assert not re.search('[\u3400-\u9fff\ufffd]',present)
assert re.findall(r'\d+(?:\.\d+)?(?:[eE][-+]?\d+)?',src)==re.findall(r'\d+(?:\.\d+)?(?:[eE][-+]?\d+)?',present)
(ROOT/'native/supplement_display_en.md').write_bytes(present.encode())
(ROOT/'native/supplement_presentation_position_log.json').write_text(json.dumps({'source_sha256':sha(src.encode()),'presentation_sha256':sha(present.encode()),'replacements':log,'reverse':'EXACT','numeric_tokens':'EXACT'},ensure_ascii=False,indent=2)+'\n', encoding='utf-8')
doc=Document();s=doc.sections[0];s.orientation=WD_ORIENT.LANDSCAPE;s.page_width=Cm(27.94);s.page_height=Cm(21.59);s.top_margin=s.bottom_margin=Cm(1.6);s.left_margin=s.right_margin=Cm(1.6)
for name,size in [('Normal',10.5),('Title',15),('Heading 1',13),('Heading 2',12),('Heading 3',11)]:
 sty=doc.styles[name];sty.font.name='Liberation Serif';sty.font.size=Pt(size);sty.font.color.rgb=RGBColor(0,0,0);sty.paragraph_format.space_after=Pt(5);sty.paragraph_format.line_spacing=1.13
 if name.startswith('Heading'):sty.paragraph_format.keep_with_next=True
 for b in list(sty.element.findall('.//'+qn('w:pBdr'))):b.getparent().remove(b)
fp=s.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER;fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),' PAGE ');fp._p.append(fld)
tables=[];panels=[];heading='';figures=[]
def ids_for(headers):
 if headers[:2]==['Benchmark','Cohort'] and 'Rule' in headers:return [0,1,headers.index('Rule')]
 if headers[:4]==['product','trend','HAC','block_length']:return [0,1,2,3]
 if headers[0]=='Run index':return [headers.index(x) for x in ['Run index','County','Product','Start','End']]
 if headers[:3]==['County','Product','Start']:return [0,1,2,3]
 if headers[0]=='County / episode':return [0,1,2]
 if headers[:3] in (['Benchmark','Counties','Rule'],['Benchmark','Cohort','Rule'],['Benchmark','Cohort','L'],['scenario','scheme','region_id']) or headers[0]=='Rate fraction':return [0,1,2]
 if headers[0]=='Date':return [0]
 if headers[0]=='Benchmark' and 'Cohort' in headers:return [0,headers.index('Cohort')]
 return [0,1]
for c in re.split(r'\n\s*\n',present.strip()):
 if c.startswith('|'):
  rows=[[v.strip() for v in l.strip().strip('|').split('|')] for l in c.splitlines() if l.startswith('|') and not re.fullmatch(r'[\s|:\-]+',l)]
  n=len(rows[0]);assert all(len(r)==n for r in rows);ti=len(tables)+1;tables.append({'table':ti,'section':heading,'rows':rows})
  ids=ids_for(rows[0]);remaining=[i for i in range(n) if i not in ids]
  groups=[list(range(n))] if n<8 else []
  if not groups:
   group=[]
   for j in remaining:
    longest=max(len(r[j]) for r in rows)
    if longest>70:
     if group:groups.append(ids+group);group=[]
     groups.append(ids+[j])
    else:
     group.append(j)
     if len(group)==max(1,7-len(ids)):groups.append(ids+group);group=[]
   if group:groups.append(ids+group)
  seen=set()
  for pn,cols in enumerate(groups,1):
   p=doc.add_paragraph(f'{heading} — table {ti}, panel {pn}/{len(groups)}');p.paragraph_format.keep_with_next=True
   if pn>1 or len(rows)>100:p.paragraph_format.page_break_before=True
   for r in p.runs:font(r,10);r.bold=True
   values=[[r[j] for j in cols] for r in rows]
   # Widths consider headers and data; dates/IDs retain sufficient recognition width.
   weights=[]
   for j in cols:
    length=max(max(len(token) for token in re.split(r'\s+',r[j]) or ['']) for r in rows)
    length=max(length,len(rows[0][j])*.65)
    weights.append(max(8,min(38,length)))
   minimum=2.0;available=24.7;extra=available-minimum*len(cols);total=sum(weights);widths=[minimum+extra*w/total for w in weights]
   t=doc.add_table(rows=len(values),cols=len(cols))
   for ri,(row,docrow) in enumerate(zip(values,t.rows)):
    for ci,(value,cell) in enumerate(zip(row,docrow.cells)):
     add_text(cell.paragraphs[0],value);seen.add((ri,cols[ci]))
   format_table(t,widths)
   for row in t.rows:
    for cell in row.cells:
     for p in cell.paragraphs:
      p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(1)
      for r in p.runs:font(r,9)
   panels.append({'source_table':ti,'panel':pn,'columns_0based':cols,'keys_0based':ids,'headers':values[0],'rows':len(values),'widths_cm':widths})
  assert len(seen)==len(rows)*n
  continue
 hm=re.match(r'^(#{1,4}) (.+)$',c)
 if hm:
  heading=hm[2];p=doc.add_paragraph(heading,'Title' if len(hm[1])==1 else 'Heading '+str(min(len(hm[1])-1,3)));continue
 fm=re.fullmatch(r'<!--\s*FIGURE:(S\d+)\s*-->',c)
 if fm:
  p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;f=ROOT/'figures'/('Figure_'+fm[1]+'.png');p.add_run().add_picture(str(f),width=Cm(22));figures.append({'file':str(f),'sha256':sha(f.read_bytes())});continue
 if c.startswith('```'):continue
 p=doc.add_paragraph();add_text(p,c);p.paragraph_format.widow_control=True
 if c.startswith('Figure '):p.paragraph_format.keep_together=True
doc.core_properties.author='';doc.core_properties.last_modified_by='';doc.core_properties.comments=''
dest=ROOT/'native/supplement.docx';doc.save(dest);clean_metadata(dest)
audit={'scientific_source_sha256':sha(src.encode()),'presentation_source_sha256':sha(present.encode()),'dictionary_sha256':sha(dp.read_bytes()),'reverse':'EXACT','source_tables':tables,'panels':panels,'figures':figures,'original_cells':sum(sum(len(r) for r in t['rows']) for t in tables),'docx_sha256':sha(dest.read_bytes())}
(ROOT/'native/supplement_structure.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2)+'\n', encoding='utf-8')
print(json.dumps({'status':'GENERATED','file':str(dest),'bytes':dest.stat().st_size,'sha256':audit['docx_sha256'],'source_tables':len(tables),'panels':len(panels),'cells':audit['original_cells'],'replacement_count':len(log)}))
